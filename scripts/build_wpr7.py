import os
import docx
from pathlib import Path

def build_wpr7():
    template_path = os.path.expanduser(r"~\Downloads\Format WPR Minor Project.docx")
    out_docx_docs = r"d:\projects\GuardShield_AI\docs\02 Weekly WPRs\WPR Week 7.docx"
    out_docx_downloads = os.path.expanduser(r"~\Downloads\GuardShield_AI_WPR_Week_7.docx")
    
    doc = docx.Document(template_path)
    
    # Update Header Paragraphs
    doc.paragraphs[2].text = 'AMITY SCHOOL OF ENGINEERING & TECHNOLOGY'
    doc.paragraphs[3].text = 'B.Tech (CSE - AIML)  VII Semester'
    doc.paragraphs[6].text = 'Project Area -- Generative AI & AI Safety (Trustworthy AI / MLOps)'
    doc.paragraphs[9].text = 'Students Weekly Progress Report (WPR — Week 7) For Odd Semester of session 2026-2027'
    
    # Fill Table 0
    t0 = doc.tables[0]
    
    t0.rows[0].cells[1].text = 'Group No. 50'
    t0.rows[1].cells[1].text = '1. Shourya Solanki\n2. Rachit Ryan Chug\n3. Dhruv Raj Singh'
    t0.rows[2].cells[1].text = '1. 01\n2. 02\n3. 03'
    t0.rows[3].cells[1].text = '1. A2305223569\n2. A2305223166\n3. A2305223191'
    t0.rows[4].cells[1].text = 'Yes. Title: "GuardShield AI: Real-Time Sidecar Proxy for LLM Hallucination Detection & Prompt Injection Defense"'
    t0.rows[5].cells[1].text = 'Yes (Approved for Group No. 50)'
    t0.rows[6].cells[1].text = 'Applied to windowed NLI factual entailment verification. Linked Cross-Encoder premise-hypothesis scoring with streaming entropy spikes.'
    t0.rows[7].cells[1].text = 'Feasible. Python 3.11+, PyTorch CUDA (RTX 3050), Cross-Encoder NLI, FastAPI async SSE proxy. Zero monetary cost (₹0).'
    t0.rows[8].cells[1].text = '₹0.00 (Fully open-source software stack and local developer environment)'
    t0.rows[9].cells[1].text = '• W1-W5: Setup, dataset curation, DeBERTa fine-tuning & evaluation (Completed)\n• W6-W8: Streaming entropy engine, NLI verification & async proxy server (In Progress / Week 7)\n• W9-W12: Comparative benchmarking vs Llama-Guard & final IEEE research paper'
    t0.rows[10].cells[1].text = '1. Windowed Cross-Encoder NLI Entailment Verifier (`nli_verifier.py`)\n2. Dual-Stage End-to-End Pipeline Integration (`guardshield_pipeline.py`)\n3. Factual Grounding & Hallucination Contradiction Scoring'
    t0.rows[11].cells[1].text = '0% (Software Project)'
    t0.rows[12].cells[1].text = '75% (Cross-Encoder NLI entailment fact-checker implemented, factual support vs contradiction scoring verified on CUDA GPU, integrated with streaming entropy triggers)'
    t0.rows[13].cells[1].text = 'Verified NLI entailment scoring on hallucinated claims vs reference premise; confirmed contradiction detection on factual drift with sub-35ms inference latency.'
    t0.rows[14].cells[1].text = 'WPR Week 7 drafted.'
    t0.rows[15].cells[1].text = '1. Shourya Solanki         2. Rachit Ryan Chug         3. Dhruv Raj Singh'
    
    # Fill Work done in this week
    p11 = doc.paragraphs[11]
    p11.text = 'Work done in this week (Week 7):\n' \
               '1. Windowed Cross-Encoder NLI Entailment Implementation: Built `src/nli_verifier.py` utilizing a Cross-Encoder architecture to classify premise-hypothesis relationships into Entailment (Factual), Neutral, or Contradiction (Hallucination).\n' \
               '2. Entropy-Triggered Grounding Loop: Integrated the NLI verifier with the streaming entropy engine, triggering asynchronous factual verification only when token uncertainty spikes past threshold (eliminating continuous compute overhead).\n' \
               '3. Latency & Contradiction Verification: Verified sub-35ms NLI cross-attention processing on local NVIDIA GeForce RTX 3050 CUDA GPU across test fact-checking cases.\n' \
               '4. End-to-End Dual-Stage Pipeline Integration: Connected Pre-Scan input filtration (DeBERTa-v3) with Post-Scan streaming entropy inspection and NLI verification into unified pipeline.'
    
    doc.save(out_docx_docs)
    doc.save(out_docx_downloads)
    print(f'Successfully built WPR 7 docx at {out_docx_docs} and {out_docx_downloads}')

if __name__ == '__main__':
    build_wpr7()
