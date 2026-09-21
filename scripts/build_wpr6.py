import os
import docx
from pathlib import Path

def build_wpr6():
    template_path = os.path.expanduser(r"~\Downloads\Format WPR Minor Project.docx")
    out_docx_docs = r"d:\projects\GuardShield_AI\docs\02 Weekly WPRs\WPR Week 6.docx"
    out_docx_downloads = os.path.expanduser(r"~\Downloads\GuardShield_AI_WPR_Week_6.docx")
    
    doc = docx.Document(template_path)
    
    # Update Header Paragraphs
    doc.paragraphs[2].text = 'AMITY SCHOOL OF ENGINEERING & TECHNOLOGY'
    doc.paragraphs[3].text = 'B.Tech (CSE - AIML)  VII Semester'
    doc.paragraphs[6].text = 'Project Area -- Generative AI & AI Safety (Trustworthy AI / MLOps)'
    doc.paragraphs[9].text = 'Students Weekly Progress Report (WPR — Week 6) For Odd Semester of session 2026-2027'
    
    # Fill Table 0
    t0 = doc.tables[0]
    
    t0.rows[0].cells[1].text = 'Group No. 50'
    t0.rows[1].cells[1].text = '1. Shourya Solanki\n2. Rachit Ryan Chug\n3. Dhruv Raj Singh'
    t0.rows[2].cells[1].text = '1. 01\n2. 02\n3. 03'
    t0.rows[3].cells[1].text = '1. A2305223569\n2. A2305223166\n3. A2305223191'
    t0.rows[4].cells[1].text = 'Yes. Title: "GuardShield AI: Real-Time Sidecar Proxy for LLM Hallucination Detection & Prompt Injection Defense"'
    t0.rows[5].cells[1].text = 'Yes (Approved for Group No. 50)'
    t0.rows[6].cells[1].text = 'Applied to streaming token entropy and async sidecar proxy. Integrated Shannon entropy math with FastAPI SSE streaming.'
    t0.rows[7].cells[1].text = 'Feasible. Python 3.11+, PyTorch CUDA, FastAPI, Uvicorn, Streaming SSE, DeBERTa-v3 SLM (86M). Zero monetary cost (₹0).'
    t0.rows[8].cells[1].text = '₹0.00 (Fully open-source software stack and local developer environment)'
    t0.rows[9].cells[1].text = '• W1-W5: Setup, dataset curation, DeBERTa fine-tuning & evaluation (Completed)\n• W6-W8: Streaming entropy engine, early termination & FastAPI sidecar proxy (In Progress / Week 6)\n• W9-W12: NLI cross-encoder integration, comparative benchmarking & final IEEE paper'
    t0.rows[10].cells[1].text = '1. Streaming Token Entropy Engine (`streaming_entropy_engine.py`)\n2. Consecutive Spike Early Termination Controller\n3. FastAPI Async Sidecar Proxy Server Prototype (`proxy_server.py`)'
    t0.rows[11].cells[1].text = '0% (Software Project)'
    t0.rows[12].cells[1].text = '65% (Real-time streaming entropy calculator implemented, early termination triggers verified on hallucinated token sequences, FastAPI proxy server prototype built)'
    t0.rows[13].cells[1].text = 'Verified sliding-window Shannon entropy H(x) on streaming logits; confirmed early stream termination on sustained spikes (>1.20 bits), preventing delivery of hallucinated tokens.'
    t0.rows[14].cells[1].text = 'WPR Week 6 drafted.'
    t0.rows[15].cells[1].text = '1. Shourya Solanki         2. Rachit Ryan Chug         3. Dhruv Raj Singh'
    
    # Fill Work done in this week
    p11 = doc.paragraphs[11]
    p11.text = 'Work done in this week (Week 6):\n' \
               '1. Streaming Token Entropy Module Implementation: Built `src/streaming_entropy_engine.py` to calculate Shannon Entropy H(x) = -sum(p * log2(p)) over discrete token probability distributions in real time.\n' \
               '2. Consecutive Spike Early Termination Logic: Designed and verified threshold-based early termination triggering when uncertainty spikes past 1.20 bits across consecutive tokens, halting hallucinated text mid-sentence.\n' \
               '3. FastAPI Sidecar Proxy Server Prototype: Developed `src/proxy_server.py` with `/v1/chat/completions` endpoint featuring asynchronous Pre-Scan security interception and Server-Sent Events (SSE) streaming token output.\n' \
               '4. Pre-Scan + Post-Scan Integration: Linked the fine-tuned DeBERTa-v3 classifier as pre-routing security middleware, returning 403 Forbidden with risk scores on detected prompt injection attempts.'
    
    doc.save(out_docx_docs)
    doc.save(out_docx_downloads)
    print(f'Successfully built WPR 6 docx at {out_docx_docs} and {out_docx_downloads}')

if __name__ == '__main__':
    build_wpr6()
