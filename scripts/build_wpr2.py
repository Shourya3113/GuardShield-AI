import docx

def build_wpr2():
    template_path = r'C:\Users\Peter\Downloads\Format WPR Minor Project.docx'
    out_docx_path = r'd:\projects\GuardShield_AI\GuardShield_AI_WPR_Week_2.docx'
    
    doc = docx.Document(template_path)
    
    # Update Header Paragraphs
    doc.paragraphs[2].text = 'AMITY SCHOOL OF ENGINEERING & TECHNOLOGY'
    doc.paragraphs[3].text = 'B.Tech (CSE - AIML)  VII Semester'
    doc.paragraphs[6].text = 'Project Area -- Generative AI & AI Safety (Trustworthy AI / MLOps)'
    doc.paragraphs[9].text = 'Students Weekly Progress Report (WPR — Week 2) For Odd Semester of session 2026-2027'
    
    # Fill Table 0
    t0 = doc.tables[0]
    
    t0.rows[0].cells[1].text = 'Group No. 50'
    t0.rows[1].cells[1].text = '1. Shourya Solanki\n2. Rachit Ryan Chug\n3. Dhruv Raj Singh'
    t0.rows[2].cells[1].text = '1. 01\n2. 02\n3. 03'
    t0.rows[3].cells[1].text = '1. A2305223569\n2. A2305223166\n3. A2305223191'
    t0.rows[4].cells[1].text = 'Yes. Title: "GuardShield AI: Real-Time Sidecar Proxy for LLM Hallucination Detection & Prompt Injection Defense"'
    t0.rows[5].cells[1].text = 'Yes (Approved for Group No. 50)'
    t0.rows[6].cells[1].text = 'Completed & applied. Formulated dataset schema mapping for 5 benchmark sources (JailbreakBench, HackAPrompt, BeaverTails, HaluEval, TruthfulQA).'
    t0.rows[7].cells[1].text = 'Feasible. Python 3.11+, PyTorch MPS/CUDA, Hugging Face Datasets pipeline, local Ollama Llama-3.1-8B. Zero monetary cost (₹0).'
    t0.rows[8].cells[1].text = '₹0.00 (Fully open-source software stack and local developer environment)'
    t0.rows[9].cells[1].text = '• W1-W3: Literature survey, setup & dataset ingestion (In Progress / Week 2)\n• W4-W7: Dataset concatenation & GuardShield-Bench-v1\n• W8-W12: DeBERTa-v3 fine-tuning & token entropy algorithm\n• W13-W16: Benchmarking & IEEE research paper\n• W17-W30 (Sem 8): FastAPI proxy, Red-Agent & Streamlit dashboard'
    t0.rows[10].cells[1].text = '1. Dataset Ingestion & Preprocessing Pipeline\n2. Pre-Scan Security Data Loader (JailbreakBench + HackAPrompt)\n3. Post-Scan Factuality Data Loader (HaluEval + TruthfulQA)\n4. Streaming Logit Entropy Extractor'
    t0.rows[11].cells[1].text = '0% (Software Project)'
    t0.rows[12].cells[1].text = '25% (Benchmark dataset ingestion, schema normalization, and streaming logit token entropy prototyping)'
    t0.rows[13].cells[1].text = 'Successfully standardized 10,000+ benchmark prompt samples into unified JSON schema; verified streaming token logit extraction with <20ms overhead.'
    t0.rows[14].cells[1].text = 'WPR Week 2 drafted.'
    t0.rows[15].cells[1].text = '1. Shourya Solanki         2. Rachit Ryan Chug         3. Dhruv Raj Singh'
    
    # Fill Work done in this week
    p11 = doc.paragraphs[11]
    p11.text = 'Work done in this week (Week 2):\n' \
               '1. Benchmark Dataset Ingestion: Ingested peer-reviewed datasets (JailbreakBench NeurIPS 2024, HackAPrompt EMNLP 2023, BeaverTails NeurIPS 2023, HaluEval EMNLP 2023).\n' \
               '2. Data Preprocessing & Normalization: Developed Python preprocessing script unifying prompt injection attacks and hallucination pairs into a standardized JSON/CSV schema.\n' \
               '3. Synthetic Data Obfuscation Design: Formulated obfuscation patterns (Base64 encoding, ROT13, Hinglish code-mixing) for synthetic `GuardShield-Bench-v1` creation.\n' \
               '4. Token Logit Entropy Prototyping: Implemented PyTorch logit entropy parser computing real-time Shannon Entropy H(x) = -sum(p * log(p)) over Ollama streaming token probabilities.\n' \
               '5. Performance Verification: Verified logit extraction overhead, achieving sub-20ms token processing latency on local MPS/CUDA GPUs.'
    
    doc.save(out_docx_path)
    print(f'Successfully built {out_docx_path}')

if __name__ == '__main__':
    build_wpr2()
